import base64
from io import BytesIO
import math
from openpyxl.styles import Font
import streamlit as st
import dashboard
import expenses
import accounts
st.set_page_config(
    page_title="Trivsys ERP",
    page_icon="💻",
    layout="wide",
)

from PIL import Image as PILImage
import pandas as pd
import payroll
import os
from openpyxl.drawing.image import Image
import shutil
from database import *
import attendance
from sales import get_filtered_sales, sales_page
from report import *
import whatsapp_integration
import purchase
import plotly.graph_objects as go
import credit_sales
from user_auth import *
import order_management
import re

# ============================================================================
# INITIALIZATION
# ============================================================================

# Ensure master database is initialized
if 'db_initialized' not in st.session_state:
    if initialize_master_db():
        st.session_state.db_initialized = True
    else:
        st.error("Failed to initialize master database")
        st.stop()

# Check for existing session
check_session()

# If a tenant session was restored, ensure tenant-local tables exist
try:
    if st.session_state.get('logged_in') and st.session_state.get('tenant'):
        ensure_tenant_tables()
except Exception:
    # Don't block startup for provisioning issues
    pass

# Set static folder where images will be stored
STATIC_FOLDER = "static"
os.makedirs(STATIC_FOLDER, exist_ok=True)

# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def clean_product_name_for_filename(product_name):
    """Convert product name to a safe filename"""
    return re.sub(r'[^a-zA-Z0-9_]', '_', product_name.strip().lower())

def save_product_image(img_data, product_name):
    """Save an RGBA image as RGB with a white background"""
    image = PILImage.open(BytesIO(img_data))
    if image.mode == "RGBA":
        white_bg = PILImage.new("RGB", image.size, (255, 255, 255))
        white_bg.paste(image, mask=image.split()[3])
        image = white_bg
    
    safe_name = clean_product_name_for_filename(product_name)
    image_filename = f"{safe_name}.jpg"
    image_path = os.path.join(STATIC_FOLDER, image_filename)
    
    os.makedirs(STATIC_FOLDER, exist_ok=True)
    image.save(image_path, "JPEG", quality=95)
    
    return f"static/{image_filename}"

def generate_excel_template():
    """Generate an Excel file template for product uploads with bold header"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Product Upload"
    
    headers = ["Product Name", "Product Category", "Selling Price", "Cost Price", "Quantity", "Image"]
    ws.append(headers)
    
    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font
    
    ws.append(["Sample Product1", "Fruits", 150.0, 100.0, 20, "sample_product1.jpg"])
    ws.append(["Sample Product2", "Fruits", 150.0, 100.0, 20, "sample_product2.jpg"])
    
    virtual_workbook = BytesIO()
    wb.save(virtual_workbook)
    virtual_workbook.seek(0)
    return virtual_workbook

def process_excel(uploaded_file, tenant):
    """Process uploaded Excel file with tenant context"""
    df = pd.read_excel(uploaded_file, engine="openpyxl")
    
    for idx, row in df.iterrows():
        name = str(row["Product Name"]).strip()
        category = str(row["Product Category"]).strip()
        price = float(row["Selling Price"])
        cost_price = float(row["Cost Price"])
        stock = int(row["Quantity"])
        if stock < 0:
            stock = 0
        
        image_path = None
        if "Image" in row and pd.notna(row["Image"]):
            original_image_filename = str(row["Image"]).strip()
            original_image_path = os.path.join(STATIC_FOLDER, original_image_filename)
            
            if os.path.exists(original_image_path):
                safe_name = clean_product_name_for_filename(name)
                new_image_filename = f"{safe_name}.jpg"
                new_image_path = os.path.join(STATIC_FOLDER, new_image_filename)
                
                if original_image_filename != new_image_filename:
                    os.rename(original_image_path, new_image_path)
                
                image_path = f"static/{new_image_filename}"
            else:
                st.warning(f"⚠️ Image '{original_image_filename}' not found for '{name}'")
        
        add_product(name, price, cost_price, stock, category, image_path)
    
    st.success("✅ Products uploaded successfully!")

def delete_product_image(product_name):
    """Delete the product image if it exists"""
    safe_name = clean_product_name_for_filename(product_name)
    image_extensions = [".jpg", ".png"]
    for ext in image_extensions:
        image_path = os.path.join(STATIC_FOLDER, f"{safe_name}{ext}")
        if os.path.exists(image_path):
            os.remove(image_path)
            return True
    return False

def get_today_adjusted_sales():
    """Return today's adjusted sales (total - returns)"""
    sales = get_filtered_sales("", "", "Daily")
    if not sales:
        return 0
    
    df = pd.DataFrame(sales, columns=[
        "SALE ID", "PRODUCTS", "QUANTITY", "TOTAL PRICE",
        "PAID AMOUNT", "DUE AMOUNT", "CUSTOMER", "DATE", "SOURCE"
    ])
    df["TOTAL PRICE"] = pd.to_numeric(df["TOTAL PRICE"], errors="coerce").fillna(0)
    total_sales = df["TOTAL PRICE"].sum()
    
    try:
        returns = fetch_returns_from_returns_table()
        total_returns = pd.DataFrame(returns)["return_amount"].sum() if returns else 0
    except Exception:
        total_returns = 0
    
    adjusted_sales = total_sales - total_returns
    return adjusted_sales

def count_low_stock_items(threshold=5):
    """Count products with stock less than or equal to threshold"""
    products = fetch_products()
    low_stock_items = [p for p in products if p["stock"] is not None and p["stock"] <= threshold]
    return len(low_stock_items)

def log_activity(activity):
    """Log an activity with timestamp and tenant info"""
    timestamp = pd.to_datetime("now").strftime("%Y-%m-%d %H:%M:%S")
    
    if st.session_state.get("logged_in", False) and "tenant" in st.session_state:
        # Use helper to safely get tenant business name
        try:
            tenant_name = get_tenant_business_name(default="Unknown")
        except Exception:
            tenant_name = "Unknown"
        log_entry = {"timestamp": timestamp, "tenant": tenant_name, "action": activity}
    else:
        log_entry = {"timestamp": timestamp, "tenant": "system", "action": activity}
    
    if "activity_log" not in st.session_state:
        st.session_state.activity_log = []
    
    st.session_state.activity_log.append(log_entry)
    
    with open("activity_log.txt", "a") as f:
        tenant_info = tenant_name if st.session_state.get("logged_in", False) else "system"
        f.write(f"{timestamp} - {tenant_info} - {activity}\n")

def get_image_base64(image_path):
    """Convert image to base64 string"""
    try:
        with open(image_path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode()
    except:
        return ""

def go_to_page(page_name):
    """Navigate to a specific page"""
    st.session_state.page = page_name
    st.rerun()

def discard_cart():
    """Clear the shopping cart"""
    st.session_state.cart.clear()
    st.session_state.total = 0.0
    st.success("Cart discarded successfully!")
    st.rerun()

def update_cart(product_id, quantity):
    """Update product quantity in the cart"""
    product = next((p for p in st.session_state.filtered_products if p["id"] == product_id), None)
    if product:
        if quantity > 0:
            if product_id in st.session_state.cart:
                st.session_state.total -= float(st.session_state.cart[product_id]["price"]) * st.session_state.cart[product_id]["quantity"]
                st.session_state.cart[product_id]["quantity"] = quantity
            else:
                st.session_state.cart[product_id] = {
                    "name": product["name"],
                    "price": float(product["price"]),
                    "quantity": quantity
                }
            st.session_state.total += float(product["price"]) * quantity
        else:
            if product_id in st.session_state.cart:
                st.session_state.total -= float(st.session_state.cart[product_id]["price"]) * st.session_state.cart[product_id]["quantity"]
                del st.session_state.cart[product_id]
        st.rerun()
    else:
        st.error(f"Product with ID {product_id} not found.")

def remove_item_from_cart(product_id):
    """Remove item from cart and update total"""
    if product_id in st.session_state.cart:
        item_price = float(st.session_state.cart[product_id]["price"]) * st.session_state.cart[product_id]["quantity"]
        st.session_state.total -= item_price
        del st.session_state.cart[product_id]
        st.success(f"Item removed from cart.")
    else:
        st.warning("Item not found in cart.")

# ============================================================================
# SESSION STATE INITIALIZATION
# ============================================================================

if "page" not in st.session_state:
    st.session_state.page = "Login"

if "cart" not in st.session_state or not isinstance(st.session_state.cart, dict):
    st.session_state.cart = {}
if "total" not in st.session_state:
    st.session_state.total = 0.0
if "products" not in st.session_state:
    st.session_state.products = []
if "filtered_products" not in st.session_state:
    st.session_state.filtered_products = []

# ============================================================================
# LOAD CSS
# ============================================================================

css_file_path = "static/style.css"
if os.path.exists(css_file_path):
    with open(css_file_path, "r") as f:
        css = f.read()
        st.markdown(f"<style>{css}</style>", unsafe_allow_html=True)

# ============================================================================
# PAGE ROUTING
# ============================================================================

# LOGIN PAGE
if st.session_state.page == "Login":
    show_login_page()

# HOME PAGE
elif st.session_state.page == "Home" and st.session_state.get("logged_in", False):
    logo_base64 = get_image_base64('static/trivsys.png')
    st.markdown(f"""
    <style>
    .container{{
        width: 100%;
        height: 50px;
        display:flex;
        align-items:center;
        }}
        .content:nth-child(1){{
           width:10%;
           height:80px;
           padding-bottom:150px;  
         }}
        .content:nth-child(1) .logo{{
            width:100%;
            height:105px;
            object-fit: cover;
        }}
        .content:nth-child(2){{
           width:90%;
           height:100%;
           padding-bottom:105px; 
           display:flex;
           justify-content:start;
         }}
         .content:nth-child(2) h3{{
            font-size:30px;
            font-style:italic;
         }}
    </style>
     <div class="container">
        <div class="content">
        <img src="data:image/png;base64,{logo_base64}" alt="Logo" class='logo'>
        </div>
        <div class='content'>
          <h3>ERP</h3>
        </div>
     </div>
""", unsafe_allow_html=True)
    # Sidebar
    with st.sidebar:
        
        # Tenant Profile Section
        tenant = get_current_tenant()
        if tenant:
            user_img = get_image_base64('static/user.png')
            business_name = get_tenant_business_name(default="Unknown")
            st.markdown(f"""
            <div style="text-align: center; padding: 0.8rem; background: rgba(255,255,255,0.1); 
                        border-radius: 20px; margin: 0.5rem 0;">
                <div style="width: 70px; height: 70px; background: linear-gradient(45deg, #fff, #fff); 
                            border-radius: 50%; margin: 0 auto 0.4rem; display: flex; align-items: center; justify-content: center;">
                    <img style= "height:50px;" src="data:image/png;base64,{user_img}" alt="user" class='user'>
                </div>
                <strong style="font-size: 1.5rem;">{business_name.upper()}</strong><br>
                <span style="background: #28a745; color: white; padding: 2px 8px; 
                            border-radius: 10px; font-size: 1rem;">
                    ACTIVE
                </span>
            </div>
            """, unsafe_allow_html=True)
        
        st.markdown("---")
        
        # Navigation Menu
        modules = {
            "Credit Sales": "Credit Sales",
            "Inventory Management": "Inventory Management",
            "Sales Module": "Sales",
            "Purchase": "Purchase",
            "POS System": "POS System",
            "Attendance Management": "Attendance Management",
            "Payroll": "Payroll",
            "Expenses":"Expenses",
            "Accounts Management":"Accounts Management",
            # "WhatsApp": "WhatsApp",
            # "Ledgers & Reports": "ledgers",
        }
        
        for icon_name, page in modules.items():
            if st.button(icon_name, key=f"nav_{page}", use_container_width=True):
                go_to_page(page)
        
        st.markdown("---")
        
        if st.button("Log out", use_container_width=True):
            logout_user()
    
    # Load products for current tenant
    if not st.session_state.products:
        st.session_state.products = fetch_products()
    
    # Quick Stats
    total_products = len(st.session_state.products)
    today_sales = get_today_adjusted_sales()
    low_stock = count_low_stock_items()
    
    col1, col2, col3 = st.columns(3)
    st.divider()
    
    st.markdown("""
        <style>
            .card {
                background-color: #fff;
                padding: 20px; 
                border-radius: 15px;
                box-shadow: 3px 3px 10px #ccc;
                text-align: left;
                border-left:5px solid #48CAE4;
            }
        </style>
    """, unsafe_allow_html=True)
    
    card_style = """
        <div class='card'>
            <strong style='font-size:20px;'>{label}</strong><br>
            <span style='font-size:24px; font-weight:bold; color:{value_color};'>
                {value}
            </span><br>
            <small><em>{note}</em></small>
        </div>
    """
    
    with col1:
        st.markdown(card_style.format(
            label="Products",
            value=f"{total_products:,}",
            note="Active inventory",
            value_color="#000000"
        ), unsafe_allow_html=True)
    
    with col2:
        st.markdown(card_style.format(
            label="Today's Sales",
            value=f"PKR {today_sales:,.0f}",
            note="Daily performance",
            value_color="#007E33"
        ), unsafe_allow_html=True)
    
    with col3:
        st.markdown(card_style.format(
            label="Stock Alerts",
            value=low_stock,
            note="Items need attention",
            value_color="#CA1B1B"
        ), unsafe_allow_html=True)
    
    # Dashboard
    try:
        dashboard.show_dashboard()
    except Exception as e:
        st.error(f"Error loading Dashboard: {str(e)}")

def create_pos_session(opening_cash):
    """
    Create a new pos_sessions row and return the new session id.
    Uses MySQL NOW() for timestamp. Status defaults to 'active'.
    """
    conn = None
    cursor = None
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        # Insert using NOW() for MySQL and %s placeholders for parameters
        cursor.execute("""
            INSERT INTO pos_sessions (opening_cash, opening_time, status)
            VALUES (%s, NOW(), 'active')
        """, (float(opening_cash),))
        conn.commit()
        session_id = cursor.lastrowid
        return session_id
    except Error as e:
        if conn:
            conn.rollback()
        logger.exception("Error creating POS session: %s", e)
        raise
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()

def close_pos_session(session_id, closing_cash):
    """Close an existing POS session"""
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            UPDATE pos_sessions 
            SET closing_time = NOW(), closing_cash = %s,status = 'closed'
            WHERE session_id = %s and status = 'active'
        """, (float(closing_cash), int(session_id)))
        conn.commit()
        log_activity(f"POS Session Closed - ID: {session_id}, Closing Cash: Rs.{closing_cash:.2f}")
        return True
    except Exception as e:
        conn.rollback()
        st.error(f"Error closing session: {e}")
        return False
    finally:
        conn.close()

def get_active_session():
    """Check if there's an active session"""
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            SELECT session_id, opening_time, opening_cash 
            FROM pos_sessions 
            WHERE closing_time IS NULL 
            ORDER BY opening_time DESC 
            LIMIT 1
        """)
        result = cursor.fetchone()
        return result
    finally:
        conn.close()
# Initialize session state variables
if "active_session_id" not in st.session_state:
    st.session_state.active_session_id = None
if "session_opening_cash" not in st.session_state:
    st.session_state.session_opening_cash = 0.0
if "flag" not in st.session_state:
    st.session_state.flag = 0

# POS SYSTEM PAGE
# POS SYSTEM PAGE
elif st.session_state.page == "POS System":
    st.title("POS System")
    
    if st.button("Back to Home"):
        go_to_page("Home")
    
    # Check for active session
    active_session = get_active_session()
    if active_session and st.session_state.active_session_id is None:
        st.session_state.active_session_id = active_session[0]
        st.session_state.session_opening_cash = active_session[2]
    # SESSION NOT STARTED - Show opening interface
    if st.session_state.active_session_id is None:
        st.markdown("### 🔓 Open POS Session")
        st.info("ℹ️ Please open a session to start using the POS system.")
        col1, col2 = st.columns([2, 1])
        with col1:
            st.markdown('##### Opening Cash Amount')
            opening_cash = st.number_input(
                "Enter Opening Cash:",
                min_value=0.0,
                value=0.0,
                step=100.0,
                format="%.2f",
                key="opening_cash_input",
                label_visibility='collapsed'
            )
        
        with col2:
            st.markdown("##### ")
            st.markdown("##### ")
            if st.button("🔓 Open Session", type="primary", use_container_width=True):
                session_id = create_pos_session(opening_cash)
                st.session_state.flag = 1
                if session_id:
                    st.session_state.active_session_id = session_id
                    st.session_state.session_opening_cash = opening_cash
                    st.success(f"✅ Session opened successfully! Session ID: {session_id}")
                    st.rerun()
                
        
        st.markdown("---")
        st.markdown("**Session Status:** 🔴 No Active Session")
    # SESSION ACTIVE - Show POS system
    st.markdown(f"**Active Session ID:** {st.session_state.active_session_id} | **Opening Cash:** Rs.{st.session_state.session_opening_cash:.2f}")
    
    # Close Session Section (at top)
    with st.expander("🔒 Close POS Session", expanded=False):
        st.warning("⚠️ Closing the session will end your current POS session.")
        
        col1, col2 = st.columns([2, 1])
        with col1:
            st.markdown('##### Closing Cash Amount')
            closing_cash = st.number_input(
                "Enter Closing Cash:",
                min_value=0.0,
                value=0.0,
                step=100.0,
                format="%.2f",
                key="closing_cash_input",
                label_visibility='collapsed'
            )
        
        with col2:
            st.markdown("##### ")
            st.markdown("##### ")
            if st.button("🔒 Close Session", type="secondary", use_container_width=True):
                if close_pos_session(st.session_state.active_session_id, closing_cash):
                    difference = float(closing_cash) - float(st.session_state.session_opening_cash)
                    st.success(f"✅ Session closed successfully!")
                    st.info(f"**Cash Difference:** Rs.{difference:.2f}")
                    
                    # Clear session data
                    st.session_state.active_session_id = None
                    st.session_state.session_opening_cash = 0.0
                    st.session_state.cart.clear()
                    st.session_state.total = 0.0
                    st.session_state.flag = 0
                    st.rerun()
                else:
                    st.error("❌ Failed to close session. Please try again.")
    if st.button("Discard Cart"):
        discard_cart()
    st.markdown("---")    
    col1, col2 = st.columns([3, 2])
    
    # LEFT COLUMN - Cart
    if st.session_state.flag:
        with col1:
            st.header("Your Cart")
            
            if st.session_state.cart:
                cart_data = [
                    {
                        "Item": f"**{item['name']}**",
                        "Price": f"**Rs.{item['price']:.2f}**",
                        "Quantity": f"**{item['quantity']}**",
                        "Subtotal": f"**Rs.{item['price'] * item['quantity']:.2f}**",
                    }
                    for item in st.session_state.cart.values()
                ]
                cart_df = pd.DataFrame(cart_data)
                st.table(cart_df)
                
                # Cart management
                st.subheader("Manage Cart Item")
                cart_product_ids = list(st.session_state.cart.keys())
                product_name_map = {pid: st.session_state.cart[pid]["name"] for pid in cart_product_ids}
                
                if cart_product_ids:
                    col1_, col2_ = st.columns([2, 1])
                    with col1_:
                        st.markdown('##### Select Product')
                        selected_pid = st.selectbox(
                            "Select Product:",
                            options=cart_product_ids,
                            format_func=lambda x: product_name_map[x],
                            label_visibility='collapsed'
                        )
                    
                    selected_item = st.session_state.cart[selected_pid]
                    product = next((p for p in st.session_state.filtered_products if p["id"] == selected_pid), None)
                    
                    if product:
                        with col2_:
                            st.markdown('##### Quantity')
                            new_quantity = st.number_input(
                                "Qty",
                                min_value=0,
                                max_value=product["stock"],
                                value=selected_item["quantity"],
                                step=1,
                                key="unified_quantity_input",
                                label_visibility='collapsed'
                            )
                        
                        if new_quantity != selected_item["quantity"]:
                            update_cart(selected_pid, new_quantity)
                            st.success(f"Quantity updated for {selected_item['name']}.")
                        
                        if st.button("Remove Selected Item"):
                            remove_item_from_cart(selected_pid)
                            st.rerun()
                
                st.write(f"**Total:** Rs.{st.session_state.total:.2f}")
                
                # Checkout section
                with st.container():
                    c1, c2 = st.columns(2)
                    c3, c4 = st.columns(2)
                    
                    with c1:
                        st.markdown('##### Enter Customer Name:')
                        customer_name = st.text_input(
                            "Enter Customer Name:",
                            key="customer_name_input",
                            max_chars=25,
                            label_visibility='collapsed'
                        )
                        if not customer_name:
                            customer_name = "Walk-in customer"
                            st.info("ℹ️ No customer entered. Defaulting to 'Walk-in customer'.")
                        customer_id = get_or_create_customer(customer_name)
                    
                    with c2:
                        st.markdown('##### Select Discount Type:')
                        discount_type = st.selectbox(
                            "Select Discount Type:",
                            ["Percentage (%)", "Amount (Rs)"],
                            key="discount_type",
                            label_visibility='collapsed'
                        )
                    
                    if discount_type == "Percentage (%)":
                        with c3:
                            st.markdown('##### Enter Discount %')
                            discount_percentage = st.number_input(
                                "Enter Discount %",
                                0, 100, 0, 1,
                                key="discount_input",
                                label_visibility='collapsed'
                            )
                        discount_amount = (st.session_state.total * discount_percentage) / 100
                    else:
                        with c3:
                            st.markdown('##### Enter Discount Amount (RS)')
                            discount_amount = st.number_input(
                                "Enter Discount Amount (Rs.)",
                                min_value=0.0,
                                value=0.0,
                                step=10.0,
                                key="flat_discount_input",
                                label_visibility='collapsed'
                            )
                    
                    final_total = max(0.0, st.session_state.total - discount_amount)
                    
                    st.markdown(f"""
                        **Discount Given:** Rs.{discount_amount:.2f}  
                        <span style='font-size:24px; font-weight:bold; color:#5a9c17;'>
                            Final Total After Discount: Rs.{final_total:.2f}
                        </span>
                    """, unsafe_allow_html=True)
                    
                    with c4:
                        st.markdown('##### Enter Paying Amount')
                        st.session_state.amount_paid = st.number_input(
                            "Enter Amount Paid:",
                            min_value=0.0,
                            format="%.2f",
                            step=10.0,
                            key="amount_paid_input_main",
                            label_visibility='collapsed'
                        )
                    
                    balance = st.session_state.amount_paid - final_total
                    if st.session_state.amount_paid > 0:
                        if balance >= 0:
                            st.success(f"✅ Balance to return: Rs.{balance:.2f}")
                        else:
                            st.warning(f"⚠️ Customer still needs to pay Rs.{abs(balance):.2f}")
                    
                    checkout_disabled = st.session_state.amount_paid == 0.0 or not customer_name
                    if st.button("Checkout✅", disabled=checkout_disabled):
                        log_activity(f"Sale Made: Final Amount = Rs.{final_total:.2f}")
                        
                        sale_cart = []
                        cart_items = []
                        stock_error = False
                        
                        for product_id, item in st.session_state.cart.items():
                            product = next((p for p in st.session_state.products if p["id"] == product_id), None)
                            if product and product["stock"] >= item["quantity"]:
                                sale_cart.append({
                                    "product_id": product_id,
                                    "quantity": item["quantity"],
                                    "sale_price": item["price"]
                                })
                                cart_items.append({
                                    "name": item["name"],
                                    "quantity": item["quantity"],
                                    "price": item["price"]
                                })
                            else:
                                st.error(f"❌ Insufficient stock for {item['name']}.")
                                stock_error = True
                        
                        if not stock_error:
                            success = record_pos_sale(sale_cart, customer_id, st.session_state.amount_paid)
                            if success:
                                for item in sale_cart:
                                    product = next((p for p in st.session_state.products if p["id"] == item["product_id"]), None)
                                    if product:
                                        update_stock(item["product_id"], product["stock"] - item["quantity"])
                                
                                st.session_state.receipt_html = generate_receipt(
                                    cart_items,
                                    final_total,
                                    float(st.session_state.amount_paid)
                                )
                                pdf_file = save_receipt_as_pdf(st.session_state.receipt_html)
                                st.session_state.cart.clear()
                                st.session_state.total = 0.0
                                st.session_state.receipt_pdf = pdf_file
                                st.session_state.products = fetch_products()  # Refresh products
                                st.success("✅ Checkout successful!")
                            else:
                                st.error("❌ Failed to record sale.")
                
                if "receipt_pdf" in st.session_state and st.session_state.receipt_pdf:
                    with open(st.session_state.receipt_pdf, "rb") as f:
                        st.download_button(
                            label="📄 Download Receipt",
                            data=f,
                            file_name="receipt.pdf",
                            mime="application/pdf"
                        )
                    os.remove(st.session_state.receipt_pdf)
                    del st.session_state.receipt_pdf
        
        # RIGHT COLUMN - Products
        with col2:
            st.header("Available Products")
            
            category_col, search_col = st.columns([1, 2])
            
            with category_col:
                st.markdown('##### Category')
                category_options = ["All"] + [c["name"] for c in fetch_categories_from_db()]
                selected_category = st.selectbox("Select Category", category_options, label_visibility="collapsed")
            
            with search_col:
                st.markdown('##### Search')
                search_query = st.text_input("🔍 Search Product by Name", label_visibility="collapsed")
            
            stock_filter = st.radio("Filter by Stock Level", ["All", "Low Stock (Less than 5)"])
            
            all_products = fetch_products(selected_category if selected_category != "All" else None)
            filtered_by_stock = [p for p in all_products if p['stock'] < 5] if stock_filter == "Low Stock (Less than 5)" else all_products
            st.session_state.filtered_products = [
                p for p in filtered_by_stock if search_query.lower() in p['name'].lower()
            ] if search_query else filtered_by_stock
            
            def image_button(image_path, label, key):
                """Simulate image button using base64"""
                if os.path.exists(image_path):
                    with open(image_path, "rb") as img_file:
                        img_bytes = img_file.read()
                        encoded = base64.b64encode(img_bytes).decode()
                    
                    custom_button = f"""
                        <button style="border: none; background: none; padding: 0;" type="submit">
                            <img src="data:image/jpeg;base64,{encoded}"
                                alt="{label}"
                                style="width:100px; height:100px; object-fit:cover; border-radius:10px;" />
                        </button>
                    """
                    st.markdown(custom_button, unsafe_allow_html=True)

                    clicked = st.button(label='Add to Cart', key=key)
                    return clicked
                else:
                    clicked = st.button(label=label, key=key)
                    return clicked
            
            if st.session_state.filtered_products:
                rows = [
                    st.session_state.filtered_products[i:i + 4]
                    for i in range(0, len(st.session_state.filtered_products), 4)
                ]
                for row in rows:
                    columns = st.columns(len(row))
                    for product, col in zip(row, columns):
                        with col:
                            st.markdown(f"**{product['name']}**")
                            image_filename = f"static/{clean_product_name_for_filename(product['name'])}.jpg"
                            image_path = image_filename if os.path.exists(image_filename) else image_filename
                            
                            clicked = image_button(image_path, product['name'], key=f"img_btn_{product['id']}")
                            if clicked:
                                if product['stock'] > 0:
                                    update_cart(product["id"], 1)
                                    st.success(f"{product['name']} added to cart!")
                                else:
                                    st.warning(f"Out of stock for {product['name']}")
                            
                            
            else:
                st.warning("No products match your filters or search.")
    else:
        st.markdown("### Open POS Session to Start Selling")
# INVENTORY MANAGEMENT PAGE
elif st.session_state.page == "Inventory Management":
    st.title("Inventory Management")
    
    if st.button("Back to Home"):
        go_to_page("Home")
    
    if 'sub_menu' not in st.session_state:
        st.session_state.sub_menu = "All Products"
    
    with st.sidebar:
        st.markdown("### Inventory Management")
        
        menu_items = [
            "All Products",
            "Add a Product",
            "Update a Product",
            "Delete a Product",
            "Add a Category",
            "Bulk Upload Products"
        ]
        
        for item in menu_items:
            if st.button(item, key=f"nav_{item}", use_container_width=True):
                st.session_state.sub_menu = item
                st.rerun()
    
    sub_menu = st.session_state.sub_menu
    
    # Bulk Upload Products
    if sub_menu == "Bulk Upload Products":
        st.header("Step 1: Upload Product Images")
        st.warning("Upload product images first. These images must match the filenames in the Excel file.")
        uploaded_images = st.file_uploader("Upload Images", type=["jpg", "jpeg", "png"], accept_multiple_files=True)
        
        if uploaded_images:
            os.makedirs(STATIC_FOLDER, exist_ok=True)
            for img in uploaded_images:
                with open(os.path.join(STATIC_FOLDER, img.name), "wb") as f:
                    f.write(img.read())
            st.success(f"✅ {len(uploaded_images)} images saved to static folder")
        
        st.divider()
        st.header("Step 2: Upload Excel File")
        
        template_file = generate_excel_template()
        st.download_button(
            label="Download Excel Template",
            data=template_file,
            file_name="product_template.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
        uploaded_file = st.file_uploader("Upload Excel File with Product Data", type=["xlsx"])
        
        if uploaded_file:
            if st.button("🚀 Process Bulk Upload"):
                tenant = get_current_tenant()
                process_excel(uploaded_file, tenant)
    
    # Update a Product
    elif sub_menu == "Update a Product":
        st.header("Update a Product")
        
        if "success_message" in st.session_state:
            st.success(st.session_state["success_message"])
            del st.session_state["success_message"]
        
        if len(st.session_state.products) == 0:
            st.warning("No products available for updating.")
        else:
            product_names = [product["name"] for product in st.session_state.products]
            selected_product_name = st.selectbox("Select Product to Update", product_names)
            
            product_to_update = next((product for product in st.session_state.products if product["name"] == selected_product_name), None)
            
            if product_to_update:
                st.subheader(f"Updating: {product_to_update['name']}")
                
                col1, col2 = st.columns(2)
                
                with col1:
                    st.markdown("##### Current Details:")
                    st.write(f"- **Name:** {product_to_update['name']}")
                    st.write(f"- **Selling Price:** Rs.{float(product_to_update['price']):.2f}")
                    st.write(f"- **Cost Price:** Rs.{float(product_to_update['cost_price']):.2f}")
                    st.write(f"- **Stock:** {product_to_update['stock']}")
                    st.write(f"- **Category:** {product_to_update['category']}")
                
                with col2:
                    image_filename = f"static/{clean_product_name_for_filename(product_to_update['name'])}.jpg"
                    current_image_path = image_filename
                    
                    if os.path.exists(current_image_path):
                        st.markdown('##### Current Product Image')
                        st.image(current_image_path, width=200)
                    else:
                        st.warning("No image found for this product.")
                
                st.markdown('---')
                col3, col4 = st.columns(2)
                col5, col6 = st.columns(2)
                
                with col3:
                    st.markdown('##### Upload New Image')
                    uploaded_image = st.file_uploader("Upload New Image", type=["jpg", "png", "jpeg"], label_visibility='collapsed')
                
                with col4:
                    st.markdown('##### New Product Name')
                    new_name = st.text_input("New Product Name", value=product_to_update["name"], label_visibility='collapsed')
                
                with col5:
                    st.markdown('##### New Selling Price')
                    new_price = st.number_input("New Selling Price", min_value=0.01, step=0.01,
                                                value=float(product_to_update["price"]) if product_to_update["price"] else 0.01,
                                                label_visibility='collapsed')
                
                with col6:
                    st.markdown('##### New Cost Price')
                    new_cost_price = st.number_input("New Cost Price", min_value=0.01, step=0.01,
                                                     value=float(product_to_update["cost_price"]) if product_to_update["cost_price"] else 0.01,
                                                     label_visibility='collapsed')
                
                col7, col8 = st.columns(2)
                
                with col7:
                    st.markdown('##### New Stock Quantity')
                    new_stock = st.number_input("New Stock Quantity", min_value=0, step=1,
                                               value=product_to_update["stock"], label_visibility='collapsed')
                
                with col8:
                    st.markdown('##### Select New Category')
                    categories = fetch_categories_from_db()
                    category_names = [category["name"] for category in categories]
                    selected_category_name = st.selectbox("Select New Category", category_names,
                                                         index=category_names.index(product_to_update["category"]) if product_to_update["category"] in category_names else 0,
                                                         label_visibility='collapsed')
                
                if st.button("Update Product"):
                    if new_name.strip():
                        old_name = product_to_update["name"].strip()
                        new_name_cleaned = new_name.strip()
                        
                        old_image_filename = clean_product_name_for_filename(old_name) + ".jpg"
                        new_image_filename = clean_product_name_for_filename(new_name_cleaned) + ".jpg"
                        
                        old_image_path = os.path.join("static", old_image_filename)
                        new_image_path = os.path.join("static", new_image_filename)
                        
                        if os.path.exists(old_image_path) and old_name != new_name_cleaned:
                            try:
                                os.rename(old_image_path, new_image_path)
                                update_image(product_to_update["id"], new_image_filename)
                            except Exception as e:
                                st.error(f"Error renaming image: {e}")
                        
                        if uploaded_image:
                            with open(new_image_path, "wb") as f:
                                shutil.copyfileobj(uploaded_image, f)
                            update_image(product_to_update["id"], new_image_filename)
                            st.success("✅ New image uploaded successfully.")
                            
                            if old_image_path != new_image_path and os.path.exists(old_image_path):
                                try:
                                    os.remove(old_image_path)
                                except Exception as e:
                                    st.warning(f"⚠️ Failed to delete old image: {e}")
                        
                        update_name(product_to_update["id"], new_name_cleaned)
                        update_price(product_to_update["id"], new_price)
                        update_cost_price(product_to_update["id"], new_cost_price)
                        update_stock(product_to_update["id"], new_stock)
                        
                        selected_category = next((category for category in categories if category["name"] == selected_category_name), None)
                        if selected_category:
                            update_category(product_to_update["id"], selected_category["id"])
                        
                        for product in st.session_state.products:
                            if product["id"] == product_to_update["id"]:
                                product["name"] = new_name_cleaned
                                product["price"] = new_price
                                product["cost_price"] = new_cost_price
                                product["stock"] = new_stock
                                product["category"] = selected_category_name
                        
                        st.session_state["success_message"] = f"""
                        ✅ **Product Updated Successfully!**  
                        - **Name:** {new_name_cleaned}  
                        - **Selling Price:** Rs. {new_price:.2f}  
                        - **Cost Price:** Rs. {new_cost_price:.2f}  
                        - **Stock:** {new_stock}  
                        - **Category:** {selected_category_name}
                        """
                        
                        st.rerun()
                    else:
                        st.error("❌ Product name cannot be empty.")
    
    # All Products
    elif sub_menu == "All Products":
        st.header("All Available Products")
        
        categories = fetch_categories_from_db()
        category_options = ["All"] + [cat["name"] for cat in categories]
        
        col1, col2, col3, col4 = st.columns([2, 3, 1, 1])
        
        with col1:
            st.markdown('#### Category')
            selected_category = st.selectbox("Category", category_options, key="category_select", label_visibility='collapsed')
        
        with col2:
            st.markdown('#### Search')
            search_query = st.text_input("Search", key="product_search", label_visibility='collapsed')
        
        all_products = fetch_products()
        
        if selected_category != "All":
            all_products = [p for p in all_products if p['category'] == selected_category]
        
        if search_query:
            all_products = [p for p in all_products if search_query.lower() in p['name'].lower()]
        
        entries_per_page = 10
        total_pages = math.ceil(len(all_products) / entries_per_page)
        
        with col3:
            st.markdown('#### Page')
            current_page = st.number_input("Page", min_value=1, max_value=total_pages if total_pages else 1,
                                          step=1, key="product_page", label_visibility='collapsed')
        
        with col4:
            st.markdown(f"**Page {int(current_page)}/{total_pages}**")
        
        start_idx = (int(current_page) - 1) * entries_per_page
        end_idx = start_idx + entries_per_page
        paginated_products = all_products[start_idx:end_idx]
        
        if paginated_products:
            st.table(pd.DataFrame(paginated_products))
        else:
            st.warning("No products match your filters or search.")
    
    # Add a Product
    elif sub_menu == "Add a Product":
        st.header("Add New Product")
        
        categories = fetch_categories_from_db()
        category_names = [category["name"] for category in categories] if categories else []
        
        col1, col2, col3, col4 = st.columns([2, 2, 2, 2])
        
        with col1:
            st.markdown('#### Product Name')
            new_product_name = st.text_input("Product Name", label_visibility='collapsed')
        
        with col2:
            st.markdown('#### Selling Price')
            new_product_price = st.number_input("Selling Price", min_value=0.01, step=0.01, label_visibility='collapsed')
        
        with col3:
            st.markdown('#### Cost Price')
            new_product_cost_price = st.number_input("Cost Price", min_value=0.01, step=0.01, label_visibility='collapsed')
        
        with col4:
            st.markdown('#### Stock Quantity')
            new_product_stock = st.number_input("Stock Quantity", min_value=0, step=1, label_visibility='collapsed')
        
        col5, col6 = st.columns([1, 1])
        
        with col5:
            st.markdown('#### Category')
            selected_category_name = st.selectbox("Category", category_names if category_names else ["General"], label_visibility='collapsed')
        
        with col6:
            st.markdown('#### Image')
            new_product_image = st.file_uploader("Image", type=["jpg", "jpeg", "png"], label_visibility='collapsed')
        
        st.markdown("---")
        
        if st.button("Add Product"):
            if new_product_name.strip():
                image_filename = None
                if new_product_image:
                    image_filename = save_product_image(new_product_image.getvalue(), new_product_name)
                
                product_id = add_product(
                    new_product_name.strip(),
                    new_product_price,
                    new_product_cost_price,
                    new_product_stock,
                    selected_category_name,
                    image_filename
                )
                
                if product_id:
                    barcode_path = generate_barcode(product_id, new_product_name.strip())
                    log_activity(f"Added new product: {new_product_name.strip()} in category '{selected_category_name}'")
                    update_product_barcode(product_id, barcode_path)
                    
                    st.session_state["success_message"] = f"✅ Product '{new_product_name.strip()}' added successfully!"
                    st.session_state.products = fetch_products()
                    st.rerun()
                else:
                    st.error("❌ Error adding product.")
            else:
                st.error("⚠️ Product name is required.")
        
        if "success_message" in st.session_state:
            st.success(st.session_state["success_message"])
            del st.session_state["success_message"]
        
        st.header("All Available Products")
        
        st.session_state.products = fetch_products()
        if st.session_state.products:
            products_df = pd.DataFrame(st.session_state.products)
            st.table(products_df)
        else:
            st.warning("No products available.")
    
    # Delete a Product
    elif sub_menu == "Delete a Product":
        st.header("Delete Product")
        
        if "products" not in st.session_state:
            st.session_state.products = fetch_products()
        
        if st.session_state.products:
            product_names = [product["name"] for product in st.session_state.products]
            st.markdown('#### Select Product to Delete')
            selected_product_name = st.selectbox("Select Product to Delete", product_names, label_visibility='collapsed')
            
            selected_product = next((product for product in st.session_state.products if product["name"] == selected_product_name), None)
            
            if st.button("Delete Product"):
                if selected_product:
                    delete_product_from_db(selected_product["id"])
                    log_activity(f"Deleted product: {selected_product_name}")
                    
                    if delete_product_image(selected_product_name):
                        st.success(f"Product '{selected_product_name}' and its image deleted successfully!")
                    else:
                        st.warning(f"Product '{selected_product_name}' deleted, but no image found.")
                    
                    st.session_state.products = [p for p in st.session_state.products if p["id"] != selected_product["id"]]
                    st.rerun()
                else:
                    st.error("Product not found!")
        
        if st.session_state.products:
            st.subheader("All Available Products")
            products_df = pd.DataFrame(st.session_state.products)
            st.table(products_df)
        else:
            st.warning("No products available to delete.")
    
    # Add a Category
    elif sub_menu == "Add a Category":
        st.header("Manage Categories")
        
        st.subheader("Add New Category")
        st.markdown('###### Category Name')
        new_category_name = st.text_input("Category Name", label_visibility='collapsed')
        
        if st.button("Add Category"):
            if new_category_name.strip():
                add_category(new_category_name.strip())
                log_activity(f"Added new category: {new_category_name.strip()}")
                st.session_state["success_message"] = f"Category '{new_category_name.strip()}' added successfully!"
                st.rerun()
            else:
                st.error("Category name cannot be empty.")
        
        if "success_message" in st.session_state:
            st.success(st.session_state["success_message"])
            del st.session_state["success_message"]
        
        st.subheader("Delete a Category")
        
        categories = fetch_categories_from_db()
        
        if categories:
            category_names = [category["name"] for category in categories]
            st.markdown('###### Select Category to Delete')
            selected_category_name = st.selectbox("Select Category to Delete", category_names, label_visibility='collapsed')
            
            selected_category = next((category for category in categories if category["name"] == selected_category_name), None)
            
            if st.button("Delete Category"):
                if selected_category:
                    delete_category(selected_category["id"])
                    log_activity(f"Deleted category: {selected_category_name}")
                    st.session_state["success_message"] = f"Category '{selected_category_name}' deleted successfully!"
                    st.rerun()
                else:
                    st.error("Category not found!")
        
        st.subheader("Existing Categories")
        
        if categories:
            category_data = {"ID": [], "Category Name": []}
            for category in categories:
                category_data["ID"].append(category["id"])
                category_data["Category Name"].append(category["name"])
            
            st.table(category_data)

# SALES MODULE PAGE
elif st.session_state.page == "Sales":
    try:
        sales_page()
    except Exception as e:
        st.error(f"Error loading Sales Module: {str(e)}")

# CREDIT SALES PAGE
elif st.session_state.page == "Credit Sales":
    try:
        credit_sales.credit_sales_page()
    except Exception as e:
        st.error(f"Error loading Credit Sales: {str(e)}")

# PURCHASE MODULE PAGE
elif st.session_state.page == "Purchase":
    try:
        purchase.purchase_page()
    except Exception as e:
        st.error(f"Error loading Purchase Module: {str(e)}")

# ATTENDANCE MANAGEMENT PAGE
elif st.session_state.page == "Attendance Management":
    try:
        attendance.attendance_page()
    except Exception as e:
        st.error(f"Error loading Attendance Management: {str(e)}")

# PAYROLL MANAGEMENT PAGE
elif st.session_state.page == "Payroll":
    try:
        payroll.payroll_page()
    except Exception as e:
        st.error(f"Error loading Payroll module: {str(e)}")

# WHATSAPP INTEGRATION PAGE
elif st.session_state.page == "WhatsApp":
    try:
        whatsapp_integration.whatsapp_page()
    except Exception as e:
        st.error(f"Error loading WhatsApp Integration: {str(e)}")
# expense module
elif st.session_state.page == "Expenses":
    try:
        expenses.expense_page()
    except Exception as e:
        st.error(f"Error loading expense module: {str(e)}")
# accounts module
elif st.session_state.page == "Accounts Management":
    try:
        accounts.accounts_page()
    except Exception as e:
        st.error(f"Error loading accounts module: {str(e)}")
# FALLBACK FOR UNKNOWN PAGES
# else:
#     st.error(f"Unknown page: {st.session_state.page}")
#     if st.button("🏠 Return to Home"):
#         go_to_page("Home")

# ENSURE USER IS LOGGED IN FOR ALL PAGES (EXCEPT LOGIN)
if st.session_state.page != "Login" and not st.session_state.get("logged_in", False):
    st.session_state.page = "Login"
    st.rerun()